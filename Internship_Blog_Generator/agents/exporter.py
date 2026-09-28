"""Exporter agent: DOCX, standalone HTML and publish-ready Markdown.

Improvements
------------
* Images are looked up in images.json (exact mapping) - no more fuzzy guessing.
* Real inline formatting (**bold**, *italic*) instead of stripping it.
* Hero image, "In this article" list, quick-facts table for each item
  (taken from the dataset), captions, page numbers.
* SEO details go into an appendix at the end, not in the middle of the article.
* HTML export embeds images (base64) + meta tags + JSON-LD schema.
* Markdown export has YAML front matter (WordPress / Hugo / Jekyll ready).
"""
from __future__ import annotations

import base64
import html
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from utils.logger import get_logger
from utils.text import clean_heading, match_item, parse_blocks, parse_inline, slugify

log = get_logger("exporter")

ACCENT = RGBColor(0x0E, 0x5E, 0x58)


def _load_json(path: Path, default):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default


class ExporterAgent:
    def __init__(self, run_dir: str | Path | None = None) -> None:
        self.run_dir = Path(run_dir) if run_dir else None

    # ------------------------------------------------------------- shared loading
    def _bundle(self, run_dir: Path | None = None) -> dict:
        rd = Path(run_dir or self.run_dir)
        blog = (rd / "blog.md").read_text(encoding="utf-8")
        return {
            "dir": rd,
            "blog": blog,
            "seo": _load_json(rd / "seo.json", {}),
            "images": _load_json(rd / "images.json", {}),
            "records": {r["name"]: r for r in _load_json(rd / "records.json", [])},
            "quality": _load_json(rd / "quality_report.json", {}),
        }

    @staticmethod
    def _image_for(heading: str, bundle: dict) -> tuple[Path | None, str]:
        info = bundle["images"].get(heading)
        if not info:
            return None, ""
        path = bundle["dir"] / "images" / info["file"]
        return (path if path.exists() else None), info.get("alt", heading)

    @staticmethod
    def _record_for(heading: str, bundle: dict) -> dict | None:
        item = match_item(heading, list(bundle["records"]))
        return bundle["records"].get(item) if item else None

    # ================================================================= DOCX
    def export_docx(self, run_dir: str | Path | None = None, output_path: str | Path | None = None) -> Path | None:
        try:
            b = self._bundle(run_dir)
        except OSError:
            log.error("blog.md not found - run the writer first.")
            return None
        doc = Document()
        self._style_document(doc)
        blocks = parse_blocks(b["blog"])
        h2s = [t for k, t in blocks if k == "h2"]
        item_h2s = [h for h in h2s if h in b["images"] and h != "hero"]

        for kind, text in blocks:
            if kind == "h1":
                title = doc.add_heading(text, level=0)
                title.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self._meta_line(doc, b["seo"])
                hero, alt = self._image_for("hero", b)
                if hero:
                    self._picture(doc, hero, alt)
                if len(item_h2s) > 2:
                    self._toc(doc, item_h2s)
            elif kind == "h2":
                doc.add_heading(text, level=1)
                img, alt = self._image_for(text, b)
                if img:
                    self._picture(doc, img, alt)
                rec = self._record_for(text, b)
                if rec and rec.get("facts"):
                    self._facts_table(doc, rec["facts"])
            elif kind == "h3":
                doc.add_heading(text, level=2)
            elif kind == "bullet":
                self._rich(doc.add_paragraph(style="List Bullet"), text)
            elif kind == "number":
                self._rich(doc.add_paragraph(style="List Number"), text)
            else:
                self._rich(doc.add_paragraph(), text)

        if b["seo"]:
            self._seo_appendix(doc, b["seo"], b["quality"])
        self._footer_page_numbers(doc)

        out = Path(output_path) if output_path else b["dir"] / "final_blog.docx"
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out)
        log.info("DOCX exported: %s", out)
        return out

    # ---- docx helpers
    @staticmethod
    def _style_document(doc: Document) -> None:
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.25
        for name, size in (("Title", 26), ("Heading 1", 18), ("Heading 2", 14)):
            st = doc.styles[name]
            st.font.name = "Calibri"
            st.font.size = Pt(size)
            st.font.color.rgb = ACCENT
        for section in doc.sections:
            section.left_margin = section.right_margin = Inches(1)

    @staticmethod
    def _rich(paragraph, text: str) -> None:
        for chunk, bold, italic in parse_inline(text):
            run = paragraph.add_run(chunk)
            run.bold, run.italic = bold or None, italic or None

    @staticmethod
    def _picture(doc, path: Path, alt: str) -> None:
        doc.add_picture(str(path), width=Inches(6.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(alt)
        run.italic, run.font.size, run.font.color.rgb = True, Pt(9), RGBColor(0x66, 0x66, 0x66)

    @staticmethod
    def _meta_line(doc, seo: dict) -> None:
        if not seo:
            return
        p = doc.add_paragraph()
        run = p.add_run(f"{seo.get('reading_time_minutes', '')} min read  •  {seo.get('word_count', '')} words")
        run.font.size, run.font.color.rgb = Pt(10), RGBColor(0x66, 0x66, 0x66)

    def _toc(self, doc, headings: list[str]) -> None:
        p = doc.add_paragraph()
        r = p.add_run("In this article")
        r.bold, r.font.color.rgb = True, ACCENT
        for i, h in enumerate(headings, 1):
            doc.add_paragraph(f"{i}. {h}").paragraph_format.space_after = Pt(0)
        doc.add_paragraph()

    @staticmethod
    def _shade(cell, hex_fill: str) -> None:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    def _facts_table(self, doc, facts: list) -> None:
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        for label, value in facts:
            row = table.add_row().cells
            row[0].text, row[1].text = str(label), str(value)
            row[0].paragraphs[0].runs[0].bold = True
            self._shade(row[0], "E6F2F1")
            row[0].width, row[1].width = Inches(1.5), Inches(4.7)
        table.columns[0].width, table.columns[1].width = Inches(1.5), Inches(4.7)
        doc.add_paragraph()

    def _seo_appendix(self, doc, seo: dict, quality: dict) -> None:
        doc.add_page_break()
        doc.add_heading("SEO & Publishing Details", level=1)
        note = doc.add_paragraph("Internal notes for the content team - remove before publishing.")
        note.runs[0].italic = True
        rows = [("SEO title", seo.get("seo_title")), ("Primary keyword", seo.get("primary_keyword")),
                ("Slug", seo.get("slug")), ("Meta description", seo.get("meta_description")),
                ("Reading time", f"{seo.get('reading_time_minutes')} minutes"),
                ("Word count", seo.get("word_count")), ("Keyword density", f"{seo.get('keyword_density_percent')}%"),
                ("SEO score", f"{seo.get('seo_score')}/100")]
        if quality:
            rows.append(("Editorial quality score", f"{quality.get('score')}/100"))
        self._facts_table(doc, [(k, v) for k, v in rows if v not in (None, "")])
        for heading, key in (("Secondary keywords", "secondary_keywords"), ("Long-tail keywords", "long_tail_keywords"),
                             ("Tags", "tags")):
            if seo.get(key):
                doc.add_heading(heading, level=2)
                for kw in seo[key]:
                    doc.add_paragraph(kw, style="List Bullet")

    @staticmethod
    def _footer_page_numbers(doc) -> None:
        para = doc.sections[0].footer.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)

    # ================================================================= HTML
    def export_html(self, run_dir: str | Path | None = None, output_path: str | Path | None = None,
                    embed_images: bool = True) -> Path | None:
        try:
            b = self._bundle(run_dir)
        except OSError:
            log.error("blog.md not found - run the writer first.")
            return None
        seo = b["seo"]
        blocks = parse_blocks(b["blog"])
        body: list[str] = []
        toc: list[tuple[str, str]] = []
        list_open = None

        def close_list():
            nonlocal list_open
            if list_open:
                body.append(f"</{list_open}>")
                list_open = None

        def inline(text: str) -> str:
            out = ""
            for chunk, bold, italic in parse_inline(text):
                chunk = html.escape(chunk)
                out += f"<strong>{chunk}</strong>" if bold else f"<em>{chunk}</em>" if italic else chunk
            return out

        def figure(heading: str) -> str:
            path, alt = self._image_for(heading, b)
            if not path:
                return ""
            src = ("data:image/png;base64," + base64.b64encode(path.read_bytes()).decode()) if embed_images \
                else f"images/{path.name}"
            return (f'<figure><img src="{src}" alt="{html.escape(alt)}" loading="lazy" width="1280" height="720">'
                    f"<figcaption>{html.escape(alt)}</figcaption></figure>")

        for kind, text in blocks:
            if kind in {"bullet", "number"}:
                tag = "ul" if kind == "bullet" else "ol"
                if list_open != tag:
                    close_list()
                    body.append(f"<{tag}>")
                    list_open = tag
                body.append(f"<li>{inline(text)}</li>")
                continue
            close_list()
            if kind == "h1":
                body.append(f"<h1>{html.escape(text)}</h1>")
                if seo:
                    body.append(f'<p class="meta">{seo.get("reading_time_minutes", "")} min read · '
                                f'{seo.get("word_count", "")} words</p>')
                body.append(figure("hero"))
                body.append("@@TOC@@")
            elif kind == "h2":
                slug = slugify(text)
                if text in b["images"] and text != "hero":
                    toc.append((slug, text))
                body.append(f'<h2 id="{slug}">{html.escape(text)}</h2>')
                body.append(figure(text))
                rec = self._record_for(text, b)
                if rec and rec.get("facts"):
                    rows = "".join(f"<tr><th>{html.escape(str(l))}</th><td>{html.escape(str(v))}</td></tr>"
                                   for l, v in rec["facts"])
                    body.append(f'<table class="facts">{rows}</table>')
            elif kind == "h3":
                body.append(f"<h3>{html.escape(text)}</h3>")
            else:
                body.append(f"<p>{inline(text)}</p>")
        close_list()

        toc_html = ""
        if len(toc) > 2:
            items = "".join(f'<li><a href="#{s}">{html.escape(t)}</a></li>' for s, t in toc)
            toc_html = f'<nav class="toc"><strong>In this article</strong><ol>{items}</ol></nav>'
        content = "\n".join(body).replace("@@TOC@@", toc_html)

        ld = "".join(f'<script type="application/ld+json">{json.dumps(o, ensure_ascii=False)}</script>'
                     for o in seo.get("json_ld", []))
        title = html.escape(seo.get("seo_title") or "Blog")
        page = HTML_TEMPLATE.format(
            title=title, description=html.escape(seo.get("meta_description", "")),
            keywords=html.escape(", ".join([seo.get("primary_keyword", ""), *seo.get("secondary_keywords", [])])),
            slug=html.escape(seo.get("slug", "")), ld=ld, content=content,
        )
        out = Path(output_path) if output_path else b["dir"] / "blog.html"
        out.write_text(page, encoding="utf-8")
        log.info("HTML exported: %s", out)
        return out

    # ================================================================= Markdown
    def export_markdown(self, run_dir: str | Path | None = None, output_path: str | Path | None = None) -> Path | None:
        try:
            b = self._bundle(run_dir)
        except OSError:
            return None
        seo = b["seo"]
        front = {
            "title": seo.get("seo_title"), "slug": seo.get("slug"), "description": seo.get("meta_description"),
            "keywords": [seo.get("primary_keyword"), *seo.get("secondary_keywords", [])],
            "tags": seo.get("tags", []), "reading_time": f"{seo.get('reading_time_minutes')} min",
        }
        lines = ["---"]
        for key, value in front.items():
            if value in (None, "", []):
                continue
            lines.append(f"{key}: {json.dumps(value, ensure_ascii=False)}")
        lines += ["---", ""]
        out = Path(output_path) if output_path else b["dir"] / "blog_publish.md"
        out.write_text("\n".join(lines) + b["blog"], encoding="utf-8")
        log.info("Markdown exported: %s", out)
        return out


HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<meta name="description" content="{description}">
<meta name="keywords" content="{keywords}">
<meta property="og:title" content="{title}">
<meta property="og:description" content="{description}">
<meta property="og:type" content="article">
<link rel="canonical" href="/{slug}">
{ld}
<style>
:root {{ --accent:#0e5e58; --bg:#fbfaf7; --ink:#1f2a2a; --muted:#667; }}
body {{ margin:0; background:var(--bg); color:var(--ink); font:18px/1.75 Georgia,'Times New Roman',serif; }}
main {{ max-width:780px; margin:0 auto; padding:32px 20px 80px; }}
h1,h2,h3 {{ font-family:'Segoe UI',Helvetica,Arial,sans-serif; line-height:1.25; color:var(--accent); }}
h1 {{ font-size:2.2rem; margin-bottom:.2em; }} h2 {{ margin-top:2.2em; border-bottom:2px solid #e3ece9; padding-bottom:.2em; }}
h3 {{ font-size:1.05rem; margin:1.4em 0 .2em; color:#234; }}
.meta {{ color:var(--muted); font:14px 'Segoe UI',sans-serif; }}
figure {{ margin:1.2em 0; }} img {{ width:100%; height:auto; border-radius:10px; }}
figcaption {{ text-align:center; font:13px 'Segoe UI',sans-serif; color:var(--muted); margin-top:.4em; }}
table.facts {{ width:100%; border-collapse:collapse; margin:1em 0; font:15px 'Segoe UI',sans-serif; }}
table.facts th {{ width:28%; text-align:left; background:#e6f2f1; }} table.facts th, table.facts td {{ padding:8px 10px; border:1px solid #d5e3e0; }}
.toc {{ background:#fff; border:1px solid #e3ece9; border-radius:10px; padding:14px 22px; font:15px 'Segoe UI',sans-serif; }}
.toc a {{ color:var(--accent); text-decoration:none; }} .toc ol {{ margin:.5em 0 0; padding-left:1.2em; }}
@media (prefers-color-scheme:dark) {{ :root {{ --bg:#141a1a; --ink:#e6ecec; }} .toc {{ background:#1b2323; border-color:#2b3737; }} h3 {{ color:#bcd; }} table.facts th {{ background:#1f3030; }} table.facts th, table.facts td {{ border-color:#2b3737; }} }}
</style>
</head>
<body><main>
{content}
</main></body>
</html>
"""
